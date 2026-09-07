"""Client Behavioral Risk Score."""

from pathlib import Path
import sqlite3

import numpy as np
import pandas as pd
import statsmodels.api as sm
from sklearn.ensemble import HistGradientBoostingClassifier
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import brier_score_loss, log_loss, precision_score, recall_score, roc_auc_score
from sklearn.model_selection import GroupShuffleSplit
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DATA_FILE = ROOT / "data" / "uci_default.xls"

PROFILE_FEATURES = ["log_limit", "age", "higher_education", "married"]
BEHAVIOR_FEATURES = [
    "latest_delay",
    "prior_late_share",
    "prior_max_delay",
    "delay_change",
    "utilization",
    "utilization_change",
    "log_mean_payment",
    "zero_payment_share",
]
ALL_FEATURES = PROFILE_FEATURES + BEHAVIOR_FEATURES

FEATURE_MART_SQL = """
DROP VIEW IF EXISTS feature_mart;

CREATE VIEW feature_mart AS
WITH ordered AS (
    SELECT
        *,
        LAG(delay_positive) OVER (
            PARTITION BY client_id
            ORDER BY month_no
        ) AS previous_delay
    FROM payment_history
),
behavior AS (
    SELECT
        client_id,
        MAX(CASE WHEN month_no = 9 THEN delay_positive END) AS latest_delay,
        AVG(
            CASE
                WHEN month_no < 9
                THEN CASE WHEN delay_positive > 0 THEN 1.0 ELSE 0.0 END
            END
        ) AS prior_late_share,
        MAX(CASE WHEN month_no < 9 THEN delay_positive END) AS prior_max_delay,
        MAX(
            CASE
                WHEN month_no = 9 THEN delay_positive - previous_delay
            END
        ) AS delay_change,
        AVG(CASE WHEN bill_amount > 0 THEN bill_amount ELSE 0 END) AS mean_bill,
        AVG(CASE WHEN month_no >= 7 THEN bill_amount END)
            - AVG(CASE WHEN month_no <= 6 THEN bill_amount END) AS bill_change,
        AVG(payment_amount) AS mean_payment,
        AVG(CASE WHEN payment_amount = 0 THEN 1.0 ELSE 0.0 END) AS zero_payment_share
    FROM ordered
    GROUP BY client_id
)
SELECT
    c.client_id,
    LN(1.0 + c.credit_limit) AS log_limit,
    c.age,
    CASE WHEN c.education IN (1, 2) THEN 1 ELSE 0 END AS higher_education,
    CASE WHEN c.marriage = 1 THEN 1 ELSE 0 END AS married,
    b.latest_delay,
    b.prior_late_share,
    b.prior_max_delay,
    b.delay_change,
    b.mean_bill / c.credit_limit AS utilization,
    b.bill_change / c.credit_limit AS utilization_change,
    LN(1.0 + b.mean_payment) AS log_mean_payment,
    b.zero_payment_share,
    o.default_flag
FROM clients AS c
JOIN behavior AS b ON c.client_id = b.client_id
JOIN outcomes AS o ON c.client_id = o.client_id;
"""


def load_source_data():
    raw = pd.read_excel(DATA_FILE, header=1)

    clients = raw[["ID", "LIMIT_BAL", "AGE", "EDUCATION", "MARRIAGE"]].copy()
    clients.columns = ["client_id", "credit_limit", "age", "education", "marriage"]

    outcomes = raw[["ID", "default payment next month"]].copy()
    outcomes.columns = ["client_id", "default_flag"]

    months = []
    status_columns = ["PAY_0", "PAY_2", "PAY_3", "PAY_4", "PAY_5", "PAY_6"]

    for i, status_column in enumerate(status_columns, start=1):
        month = raw[["ID", status_column, f"BILL_AMT{i}", f"PAY_AMT{i}"]].copy()
        month.columns = ["client_id", "payment_status", "bill_amount", "payment_amount"]
        month["month_no"] = 10 - i
        month["delay_positive"] = month["payment_status"].clip(lower=0)
        months.append(month)

    history = pd.concat(months, ignore_index=True).sort_values(["client_id", "month_no"])
    return raw, clients, outcomes, history


def build_feature_mart(clients, outcomes, history):
    with sqlite3.connect(":memory:") as connection:
        connection.create_function("LN", 1, np.log)
        clients.to_sql("clients", connection, if_exists="replace", index=False)
        outcomes.to_sql("outcomes", connection, if_exists="replace", index=False)
        history.to_sql("payment_history", connection, if_exists="replace", index=False)
        connection.executescript(FEATURE_MART_SQL)
        mart = pd.read_sql_query("SELECT * FROM feature_mart ORDER BY client_id", connection)
    return mart


def make_splits(raw, mart):
    y = mart["default_flag"]
    groups = pd.util.hash_pandas_object(raw.iloc[:, 1:-1], index=False)

    train_validation, test = next(
        GroupShuffleSplit(n_splits=1, test_size=0.20, random_state=42).split(mart, y, groups)
    )

    train_relative, validation_relative = next(
        GroupShuffleSplit(n_splits=1, test_size=0.25, random_state=43).split(
            mart.iloc[train_validation],
            y.iloc[train_validation],
            groups.iloc[train_validation],
        )
    )

    train = train_validation[train_relative]
    validation = train_validation[validation_relative]

    split = pd.Series("train", index=mart.index)
    split.iloc[validation] = "validation"
    split.iloc[test] = "test"
    mart = mart.copy()
    mart["split"] = split

    return mart, train, validation, test, groups


def logistic_model():
    return make_pipeline(
        StandardScaler(),
        LogisticRegression(C=np.inf, solver="lbfgs", max_iter=2000, tol=1e-9),
    )


def train_models(mart, train, validation, test):
    y = mart["default_flag"]
    models = {
        "Profile": (logistic_model(), PROFILE_FEATURES),
        "Full": (logistic_model(), ALL_FEATURES),
        "Boosting": (
            HistGradientBoostingClassifier(
                max_iter=150,
                max_leaf_nodes=15,
                learning_rate=0.05,
                l2_regularization=1,
                early_stopping=False,
                random_state=42,
            ),
            ALL_FEATURES,
        ),
    }

    metrics = []

    for model_name, (model, features) in models.items():
        model.fit(mart.loc[train, features], y.iloc[train])

        for split_name, rows in (("validation", validation), ("test", test)):
            probabilities = model.predict_proba(mart.loc[rows, features])[:, 1]
            metrics.append(
                {
                    "model": model_name,
                    "split": split_name,
                    "n": len(rows),
                    "roc_auc": roc_auc_score(y.iloc[rows], probabilities),
                    "brier": brier_score_loss(y.iloc[rows], probabilities),
                    "log_loss": log_loss(y.iloc[rows], probabilities),
                    "precision": precision_score(y.iloc[rows], probabilities >= 0.5, zero_division=0),
                    "recall": recall_score(y.iloc[rows], probabilities >= 0.5, zero_division=0),
                }
            )

    return models, pd.DataFrame(metrics)


def score_clients(mart, full_model, train):
    mart = mart.copy()
    mart["pd"] = full_model.predict_proba(mart[ALL_FEATURES])[:, 1]
    mart["risk_score"] = 100 * mart["pd"]

    cuts = np.quantile(mart.loc[train, "pd"], [0.2, 0.4, 0.6, 0.8])
    mart["risk_band"] = np.array(list("ABCDE"))[
        np.searchsorted(cuts, mart["pd"], side="left")
    ]
    return mart, cuts


def estimate_coefficients(mart, full_model, train):
    y = mart["default_flag"]
    scaler = full_model.named_steps["standardscaler"]
    x = pd.DataFrame(scaler.transform(mart[ALL_FEATURES]), columns=ALL_FEATURES)

    fit = sm.Logit(y.iloc[train], sm.add_constant(x.iloc[train])).fit(disp=False, maxiter=100)
    ci = fit.conf_int()

    coefficients = pd.DataFrame(
        {
            "feature": ALL_FEATURES,
            "coef": fit.params.iloc[1:].values / scaler.scale_,
            "std_error": fit.bse.iloc[1:].values / scaler.scale_,
            "p_value": fit.pvalues.iloc[1:].values,
            "ci_low": ci.iloc[1:, 0].values / scaler.scale_,
            "ci_high": ci.iloc[1:, 1].values / scaler.scale_,
        }
    )

    logistic = full_model.named_steps["logisticregression"]
    weights = logistic.coef_[0] / scaler.scale_
    intercept = float(logistic.intercept_[0] - weights @ scaler.mean_)

    return coefficients, intercept, weights


def review_capacity(mart, test):
    ranked = mart.loc[test].sort_values(["pd", "client_id"], ascending=[False, True])
    total_defaults = ranked["default_flag"].sum()
    base_rate = ranked["default_flag"].mean()

    rows = []
    for fraction in (0.05, 0.10, 0.20):
        selected = ranked.head(int(np.ceil(fraction * len(ranked))))
        precision = selected["default_flag"].mean()
        rows.append(
            {
                "capacity": fraction,
                "reviewed": len(selected),
                "defaults": int(selected["default_flag"].sum()),
                "precision": precision,
                "recall": selected["default_flag"].sum() / total_defaults,
                "lift": precision / base_rate,
            }
        )
    return pd.DataFrame(rows)


def build_final_excel(mart, metrics, queue):
    final_dir = ROOT / "final"
    final_dir.mkdir(exist_ok=True)
    output_file = final_dir / "Client_Behavioral_Risk_Score.xlsx"

    wb = Workbook()

    ws = wb.active
    ws.title = "Portfolio"
    ws["A1"] = "Client Behavioral Risk Score"
    ws["A1"].font = Font(size=18, bold=True)
    ws["A3"] = "Test clients"
    ws["B3"] = int((mart["split"] == "test").sum())
    ws["D3"] = "Defaults"
    ws["E3"] = int(mart.loc[mart["split"] == "test", "default_flag"].sum())
    ws["A4"] = "Default rate"
    ws["B4"] = float(mart.loc[mart["split"] == "test", "default_flag"].mean())
    ws["D4"] = "Mean risk score"
    ws["E4"] = float(mart.loc[mart["split"] == "test", "risk_score"].mean())
    ws["B4"].number_format = "0.00%"
    ws["E4"].number_format = "0.00"

    for cell in ("A3", "D3", "A4", "D4"):
        ws[cell].font = Font(bold=True)

    bands = (
        mart.query("split == 'test'")
        .groupby("risk_band", as_index=False)
        .agg(
            clients=("client_id", "size"),
            defaults=("default_flag", "sum"),
            default_rate=("default_flag", "mean"),
            mean_pd=("pd", "mean"),
        )
        .sort_values("risk_band")
    )

    start_row = 7
    headers = ["Risk band", "Clients", "Defaults", "Default rate", "Mean PD"]
    for col, value in enumerate(headers, 1):
        ws.cell(start_row, col, value).font = Font(bold=True)

    for r, row in enumerate(bands.itertuples(index=False), start_row + 1):
        ws.cell(r, 1, row.risk_band)
        ws.cell(r, 2, int(row.clients))
        ws.cell(r, 3, int(row.defaults))
        ws.cell(r, 4, float(row.default_rate)).number_format = "0.00%"
        ws.cell(r, 5, float(row.mean_pd)).number_format = "0.00%"

    chart = LineChart()
    chart.title = "Observed default rate by risk band"
    chart.y_axis.title = "Default rate"
    chart.x_axis.title = "Risk band"
    data = Reference(ws, min_col=4, min_row=start_row, max_row=start_row + len(bands))
    cats = Reference(ws, min_col=1, min_row=start_row + 1, max_row=start_row + len(bands))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 8
    chart.width = 14
    ws.add_chart(chart, "G7")

    quality = wb.create_sheet("Model Quality")
    quality["A1"] = "Model Quality"
    quality["A1"].font = Font(size=18, bold=True)

    test_metrics = metrics.query("split == 'test'").reset_index(drop=True)
    metric_headers = ["Model", "ROC-AUC", "Brier", "Log Loss", "Precision", "Recall"]
    for col, value in enumerate(metric_headers, 1):
        quality.cell(3, col, value).font = Font(bold=True)

    for r, row in enumerate(test_metrics.itertuples(index=False), 4):
        quality.cell(r, 1, row.model)
        quality.cell(r, 2, float(row.roc_auc))
        quality.cell(r, 3, float(row.brier))
        quality.cell(r, 4, float(row.log_loss))
        quality.cell(r, 5, float(row.precision))
        quality.cell(r, 6, float(row.recall))
        for c in range(2, 7):
            quality.cell(r, c).number_format = "0.0000"

    auc_chart = BarChart()
    auc_chart.title = "ROC-AUC on test"
    auc_chart.y_axis.title = "ROC-AUC"
    auc_data = Reference(quality, min_col=2, min_row=3, max_row=3 + len(test_metrics))
    auc_cats = Reference(quality, min_col=1, min_row=4, max_row=3 + len(test_metrics))
    auc_chart.add_data(auc_data, titles_from_data=True)
    auc_chart.set_categories(auc_cats)
    auc_chart.height = 8
    auc_chart.width = 14
    quality.add_chart(auc_chart, "H3")

    queue_start = 10
    quality.cell(queue_start, 1, "Review capacity").font = Font(bold=True)
    queue_headers = ["Capacity", "Reviewed", "Defaults", "Precision", "Recall", "Lift"]
    for col, value in enumerate(queue_headers, 1):
        quality.cell(queue_start + 1, col, value).font = Font(bold=True)

    for r, row in enumerate(queue.itertuples(index=False), queue_start + 2):
        quality.cell(r, 1, float(row.capacity)).number_format = "0%"
        quality.cell(r, 2, int(row.reviewed))
        quality.cell(r, 3, int(row.defaults))
        quality.cell(r, 4, float(row.precision)).number_format = "0.00%"
        quality.cell(r, 5, float(row.recall)).number_format = "0.00%"
        quality.cell(r, 6, float(row.lift)).number_format = "0.00"

    explorer = wb.create_sheet("Client Explorer")
    export_columns = [
        "client_id", "age", "higher_education", "married",
        "latest_delay", "prior_late_share", "prior_max_delay", "delay_change",
        "utilization", "utilization_change", "log_mean_payment", "zero_payment_share",
        "default_flag", "split", "pd", "risk_score", "risk_band"
    ]
    for col, value in enumerate(export_columns, 1):
        cell = explorer.cell(1, col, value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center")

    for row_index, values in enumerate(mart[export_columns].itertuples(index=False, name=None), 2):
        for col_index, value in enumerate(values, 1):
            explorer.cell(row_index, col_index, value)

    explorer.freeze_panes = "A2"
    explorer.auto_filter.ref = explorer.dimensions

    for ws_ in (ws, quality, explorer):
        for column_cells in ws_.columns:
            max_length = min(max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells[:200]), 28)
            ws_.column_dimensions[get_column_letter(column_cells[0].column)].width = max(10, max_length + 2)

    wb.save(output_file)
    return output_file


def verify(raw, history, mart, groups, intercept, weights):
    assert len(raw) == 30_000
    assert len(mart) == 30_000
    assert mart["client_id"].is_unique
    assert len(history) == 180_000
    assert history.groupby("client_id").size().eq(6).all()
    assert mart["risk_score"].between(0, 100).all()
    assert mart["default_flag"].sum() == 6_636
    assert mart.query("split == 'test'").shape[0] == 6_000
    assert not mart.isna().any().any()

    group_check = pd.DataFrame({"client_id": raw["ID"], "group_id": groups}).merge(
        mart[["client_id", "split"]]
    )
    assert group_check.groupby("group_id")["split"].nunique().max() == 1

    linear = intercept + mart[ALL_FEATURES].to_numpy() @ weights
    reconstructed = 100 / (1 + np.exp(-linear))
    assert np.allclose(reconstructed, mart["risk_score"], atol=1e-7)


def main():
    raw, clients, outcomes, history = load_source_data()
    mart = build_feature_mart(clients, outcomes, history)
    mart, train, validation, test, groups = make_splits(raw, mart)

    models, metrics = train_models(mart, train, validation, test)
    mart, _ = score_clients(mart, models["Full"][0], train)
    coefficients, intercept, weights = estimate_coefficients(mart, models["Full"][0], train)
    queue = review_capacity(mart, test)

    verify(raw, history, mart, groups, intercept, weights)
    final_file = build_final_excel(mart, metrics, queue)

    print("\nMODEL QUALITY")
    print(metrics.round(4).to_string(index=False))

    print("\nTOP-RISK REVIEW")
    print(queue.round(4).to_string(index=False))

    print("\nLOGISTIC REGRESSION COEFFICIENTS")
    print(coefficients.round(4).to_string(index=False))

    print(f"\nFinal Excel: {final_file}")
    print("PASS: data, split and score checks completed successfully.")


if __name__ == "__main__":
    main()
